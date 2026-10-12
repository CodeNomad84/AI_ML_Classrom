"""Convert a slide outline (Markdown) into a .pptx deck with python-pptx.

Provenance: new content (course tooling).

Outline format (one block per slide)::

    ## SLIDE 3 | Slide title
    **Content:**
    - bullet (max 5)
    **Visual:** IMAGE ../assets/week01/figure.png   or a text description
    **Notes:** speaker notes

A YAML-like front matter (``title``, ``subtitle``, ``course``, ``footer``,
``language``, ``provenance``) produces the cover slide. ``language: fa`` lays
the slide out right-to-left and uses a Persian-capable font.

Usage:
    python tools/build_slides.py slides/en/week01_python_environment_EN.md [-o out.pptx]
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Final

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDE_WIDTH: Final = Inches(13.333)
SLIDE_HEIGHT: Final = Inches(7.5)
MARGIN: Final = Inches(0.5)
TITLE_TOP: Final = Inches(0.35)
TITLE_HEIGHT: Final = Inches(0.95)
BODY_TOP: Final = Inches(1.5)
BODY_HEIGHT: Final = Inches(5.3)
COLUMN_GAP: Final = Inches(0.3)
TEXT_COLUMN_WIDTH: Final = Inches(6.4)
FOOTER_TOP: Final = Inches(7.0)
MAX_BULLETS: Final[int] = 5

NAVY: Final = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT: Final = RGBColor(0xE0, 0x7A, 0x1F)
GREY: Final = RGBColor(0x59, 0x59, 0x59)
LIGHT: Final = RGBColor(0xF2, 0xF2, 0xF2)

FONTS: Final[dict[str, str]] = {"en": "Calibri", "fa": "Tahoma"}
LANG_TAGS: Final[dict[str, str]] = {"en": "en-US", "fa": "fa-IR"}
EXAMPLE_PREFIXES: Final[tuple[str, ...]] = ("Example", "مثال")
SLIDE_HEADER: Final = re.compile(r"^## SLIDE (\d+) \| (.+)$", re.M)
FIELD: Final = re.compile(r"^\*\*(Content|Visual|Notes):\*\*\s*(.*)$")


@dataclass
class SlideSpec:
    """One slide parsed from the outline."""

    number: int
    title: str
    bullets: list[str] = field(default_factory=list)
    visual: str = ""
    notes: str = ""

    @property
    def is_example(self) -> bool:
        """Return True for worked-example slides (highlighted in the deck)."""
        return self.title.startswith(EXAMPLE_PREFIXES)

    @property
    def image(self) -> str | None:
        """Return the relative image path if the visual is ``IMAGE <path>``."""
        match = re.match(r"IMAGE\s+(\S+)", self.visual)
        return match.group(1) if match else None


@dataclass
class Outline:
    """Front matter plus slides of one outline file."""

    meta: dict[str, str]
    slides: list[SlideSpec]

    @property
    def language(self) -> str:
        """Return the deck language code ("en" or "fa")."""
        return self.meta.get("language", "en")


def parse_outline(text: str) -> Outline:
    """Parse the outline Markdown into an :class:`Outline`.

    Args:
        text: Full Markdown source.

    Returns:
        Parsed outline.

    Raises:
        ValueError: If a slide has more than MAX_BULLETS bullets.
    """
    meta: dict[str, str] = {}
    front = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if front:
        for line in front.group(1).splitlines():
            key, _, value = line.partition(":")
            meta[key.strip()] = value.strip().strip('"')

    slides: list[SlideSpec] = []
    headers = list(SLIDE_HEADER.finditer(text))
    for i, header in enumerate(headers):
        end = headers[i + 1].start() if i + 1 < len(headers) else len(text)
        spec = SlideSpec(int(header.group(1)), header.group(2).strip())
        current = ""
        for raw in text[header.end() : end].splitlines():
            line = raw.strip()
            match = FIELD.match(line)
            if match:
                current, rest = match.group(1), match.group(2)
                if current == "Visual":
                    spec.visual = rest
                elif current == "Notes":
                    spec.notes = rest
            elif current == "Content" and line.startswith("- "):
                spec.bullets.append(line[2:].strip())
            elif (
                current in ("Visual", "Notes") and line and not line.startswith("<!--")
            ):
                attr = current.lower()
                setattr(spec, attr, f"{getattr(spec, attr)} {line}".strip())
        if len(spec.bullets) > MAX_BULLETS:
            raise ValueError(
                f"Slide {spec.number} has {len(spec.bullets)} bullets (max 5)"
            )
        slides.append(spec)
    return Outline(meta, slides)


class DeckBuilder:
    """Render an :class:`Outline` into a 16:9 python-pptx presentation."""

    def __init__(self, outline: Outline, base_dir: Path) -> None:
        """Prepare an empty widescreen presentation.

        Args:
            outline: Parsed outline.
            base_dir: Folder used to resolve relative image paths.
        """
        self.outline = outline
        self.base_dir = base_dir
        self.rtl = outline.language == "fa"
        self.font = FONTS.get(outline.language, FONTS["en"])
        self.lang_tag = LANG_TAGS.get(outline.language, LANG_TAGS["en"])
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_WIDTH
        self.prs.slide_height = SLIDE_HEIGHT
        self._blank = self.prs.slide_layouts[6]

    def build(self) -> Presentation:
        """Create the cover slide and one slide per outline entry."""
        self._cover()
        for spec in self.outline.slides:
            self._content_slide(spec)
        return self.prs

    def _style_paragraph(
        self,
        paragraph,
        size: int,
        bold: bool = False,
        color: RGBColor = NAVY,
        italic: bool = False,
        align: PP_ALIGN | None = None,
    ) -> None:
        """Apply font, direction and alignment to every run of ``paragraph``."""
        paragraph.alignment = align or (PP_ALIGN.RIGHT if self.rtl else PP_ALIGN.LEFT)
        if self.rtl:
            paragraph._p.get_or_add_pPr().set("rtl", "1")
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.name = self.font
            run.font.color.rgb = color
            rpr = run._r.get_or_add_rPr()
            rpr.set("lang", self.lang_tag)
            for tag in ("a:cs", "a:ea"):
                element = rpr.find(qn(tag))
                if element is None:
                    element = rpr.makeelement(qn(tag), {})
                    rpr.append(element)
                element.set("typeface", self.font)

    def _text_box(
        self,
        slide,
        left: Emu,
        top: Emu,
        width: Emu,
        height: Emu,
        lines: list[str],
        size: int,
        **style,
    ) -> None:
        """Add a word-wrapped text box with one paragraph per line."""
        frame = slide.shapes.add_textbox(left, top, width, height).text_frame
        frame.word_wrap = True
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.space_after = Pt(10)
            self._style_paragraph(paragraph, size, **style)

    def _cover(self) -> None:
        """Cover slide with title, subtitle, course and provenance."""
        meta = self.outline.meta
        slide = self.prs.slides.add_slide(self._blank)
        band = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(2.2), SLIDE_WIDTH, Inches(2.4)
        )
        band.fill.solid()
        band.fill.fore_color.rgb = NAVY
        band.line.fill.background()
        width = SLIDE_WIDTH - 2 * MARGIN
        self._text_box(
            slide,
            MARGIN,
            Inches(2.4),
            width,
            Inches(1.1),
            [meta.get("title", "")],
            40,
            bold=True,
            color=RGBColor(255, 255, 255),
        )
        self._text_box(
            slide,
            MARGIN,
            Inches(3.5),
            width,
            Inches(0.8),
            [meta.get("subtitle", "")],
            24,
            color=RGBColor(255, 255, 255),
        )
        self._text_box(
            slide,
            MARGIN,
            Inches(5.0),
            width,
            Inches(0.6),
            [meta.get("course", "")],
            20,
            color=NAVY,
        )
        self._text_box(
            slide,
            MARGIN,
            Inches(6.2),
            width,
            Inches(0.9),
            [meta.get("provenance", "")],
            11,
            color=GREY,
            italic=True,
        )

    def _content_slide(self, spec: SlideSpec) -> None:
        """Title, bullets on one side, visual on the other, footer and notes."""
        slide = self.prs.slides.add_slide(self._blank)
        width = SLIDE_WIDTH - 2 * MARGIN
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.15)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT if spec.is_example else NAVY
        bar.line.fill.background()
        self._text_box(
            slide,
            MARGIN,
            TITLE_TOP,
            width,
            TITLE_HEIGHT,
            [spec.title],
            30,
            bold=True,
            color=ACCENT if spec.is_example else NAVY,
        )

        visual_width = width - TEXT_COLUMN_WIDTH - COLUMN_GAP
        if self.rtl:
            text_left, visual_left = MARGIN + visual_width + COLUMN_GAP, MARGIN
        else:
            text_left, visual_left = MARGIN, MARGIN + TEXT_COLUMN_WIDTH + COLUMN_GAP
        self._text_box(
            slide,
            text_left,
            BODY_TOP,
            TEXT_COLUMN_WIDTH,
            BODY_HEIGHT,
            [f"• {b}" for b in spec.bullets],
            20,
            color=RGBColor(0x26, 0x26, 0x26),
        )
        self._visual(slide, spec, visual_left, visual_width)

        footer = f"{self.outline.meta.get('footer', '')}   |   {spec.number}"
        self._text_box(
            slide, MARGIN, FOOTER_TOP, width, Inches(0.4), [footer], 11, color=GREY
        )
        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    def _visual(self, slide, spec: SlideSpec, left: Emu, width: Emu) -> None:
        """Insert the image (scaled to fit) or a labelled placeholder box."""
        image = spec.image
        if image and (self.base_dir / image).exists():
            picture = slide.shapes.add_picture(
                str(self.base_dir / image), left, BODY_TOP
            )
            scale = min(width / picture.width, BODY_HEIGHT / picture.height)
            picture.width = int(picture.width * scale)
            picture.height = int(picture.height * scale)
            picture.left = int(left + (width - picture.width) / 2)
            return
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, BODY_TOP, width, BODY_HEIGHT
        )
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT
        box.line.color.rgb = GREY
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        label = "تصویر: " if self.rtl else "Visual: "
        frame.paragraphs[0].text = label + spec.visual.replace("`", "")
        self._style_paragraph(
            frame.paragraphs[0], 16, color=GREY, italic=True, align=PP_ALIGN.CENTER
        )


def main() -> None:
    """Command-line entry point."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("outline", type=Path)
    parser.add_argument("-o", "--output", type=Path, default=None)
    args = parser.parse_args()
    outline = parse_outline(args.outline.read_text(encoding="utf-8"))
    output = args.output or args.outline.with_suffix(".pptx")
    DeckBuilder(outline, args.outline.parent).build().save(output)
    print(f"wrote {output} ({len(outline.slides) + 1} slides)")


if __name__ == "__main__":
    main()
